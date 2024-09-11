from utils.config.open_ai import PROMPTS,OPENAI
from openai import OpenAI
import re
from openpyxl.drawing.image import Image
import pandas as pd
from utils.constants import SYSTEM_MESSAGE ,ADHERENCE_SCORE_NOT_FOUND , RELEVANCE_SCORE_NOT_FOUND, DATAFRAME_EMPTY
from utils.graphs.graphs import generate_backlog_duration_graph, generate_epic_graph, generate_issue_type_graph, generate_ticket_scores_graph


# OpenAI setup
client = OpenAI(api_key=OPENAI['api_key'])
model = OPENAI['model']

#Relevance Score
def find_relevance_score(description_to_check, summary,placeholders):
    prompt = PROMPTS['relevance'].format(description=description_to_check, summary=summary, placeholders=placeholders)
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": SYSTEM_MESSAGE},
            {"role": "user", "content": prompt}
        ]
    ) 
    content = response.choices[0].message.content
    match = re.search(r"Relevance Score\s*:\s*(\d+)", content, re.IGNORECASE)
    if match:
        relevance_score = float(match.group(1))
    else:
        print(RELEVANCE_SCORE_NOT_FOUND)
        relevance_score = 0.0  # Default or error value if score not found
    return relevance_score * 10  # Convert to percentage

#Adherence Score
def find_adherence_score(description_to_check, headings, placeholders):

    joined_headings = ", ".join(headings)
    prompt = PROMPTS['adherence'].format(description=description_to_check, headings=joined_headings)
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": SYSTEM_MESSAGE},
            {"role": "user", "content": prompt}
        ]
    )
    content = response.choices[0].message.content
    match = re.search(r"Adherence Score:\s*(0(\.\d+)?|1(\.0+)?)", content)
    if match:
        adherence_score = float(match.group(1))
    else:
        print(ADHERENCE_SCORE_NOT_FOUND)
        adherence_score = 0.0  # Default or error value if score not found
    return adherence_score

#Generate New description for tickets which do not a 100% adherence score
def generate_perfect_adherence_description(description_to_check, headings):

    joined_headings = ", ".join(headings)
    prompt = PROMPTS['adherence_restructure'].format(description=description_to_check, headings=joined_headings)
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": SYSTEM_MESSAGE},
            {"role": "user", "content": prompt}
        ]
    )
    content = response.choices[0].message.content
    return content.strip()

# Generate backlog report with graph embedded

def generate_backlog_report(active_workbook, workbook, excel_file_path):
    
    # Convert the worksheet to a DataFrame for processing
    data = active_workbook.values
    columns = next(data)[0:]  # Get the header
    df = pd.DataFrame(data, columns=columns)

    if df.empty:
        print(DATAFRAME_EMPTY)
        return

    # Generate the graphs
    issue_type_graph_buffer = generate_issue_type_graph(df)  # Correct function call
    backlog_graph_buffer = generate_backlog_duration_graph(df, days_threshold=20)
    epic_graph_buffer = generate_epic_graph(df)  # Correct function call

    # Load and add the images from the buffers to the Excel sheet
    img_backlog = Image(backlog_graph_buffer)  # Convert backlog graph buffer to Image
    active_workbook.add_image(img_backlog, 'E2')  # Position the backlog graph

    img_epic = Image(epic_graph_buffer)  # Convert epic graph buffer to Image
    active_workbook.add_image(img_epic, 'E20')  # Position the epic graph

    img_issue_type = Image(issue_type_graph_buffer)  # Convert issue type graph buffer to Image
    active_workbook.add_image(img_issue_type, 'E38')  # Position the issue type graph

    # Save the Excel file
    workbook.save(excel_file_path)


# Generate Sprint report with graph embedded 
def generate_sprint_report(active_workbook,workbook, excel_file_path):

    # Convert the worksheet to a DataFrame for plotting
    df = pd.DataFrame(active_workbook.values)

    if df.empty:
        print(DATAFRAME_EMPTY)
    else:
        df.columns = df.iloc[0]  # Set the first row as the header
        df = df[1:]  # Remove the first row from the data

        # Convert the score columns to numeric (float) data types
        df['Relevance Score (%)'] = pd.to_numeric(df['Relevance Score (%)'])
        df['Adherence Score (%)'] = pd.to_numeric(df['Adherence Score (%)'])
        df['Total Score (%)'] = pd.to_numeric(df['Total Score (%)'])

        # Generate a graph from the DataFrame
        sprint_graph_buffer = generate_ticket_scores_graph(df)

        # Load and add the image from the buffer to the Excel sheet
        img = Image(sprint_graph_buffer)
        active_workbook.add_image(img, 'M2')  # Position the image in the Excel sheet

    # Save the Excel file
    workbook.save(excel_file_path)
